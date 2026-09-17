"""
excel_rule_parser.py

A deterministic reader for business rules encoded in an Excel
template's own structure rather than in prose — dropdown/range/length
data validation, cross-column formulas ("Total = Qty * Price"),
conditional formatting, and (for .xlsm) VBA macro source.

This exists alongside validation_service.infer_column_rules the same
way sql_constraint_parser.py exists alongside the Groq-based rules
parser: a deterministic safety net that reads the source directly
rather than guessing, for the layers where the source is unambiguous
(data validation, formulas). Conditional formatting and VBA are read
too, but returned separately as advisory material rather than folded
into column_rules, since a highlight color or an arbitrary macro is
not confirmed business logic the way a dropdown's allowed values are.

Deliberately conservative: a validation type, operator, or formula
shape this module doesn't recognise is skipped rather than guessed at.
"""
import logging
import re

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from app.services import safe_expr

logger = logging.getLogger(__name__)

try:
    from oletools.olevba import VBA_Parser
except Exception:  # pragma: no cover - optional dependency
    VBA_Parser = None


# ============================================================
# SHARED HELPERS
# ============================================================

def _header_map(ws) -> dict:
    """{column_index (1-based): header name} read from row 1."""
    header = {}
    first_row = next(ws.iter_rows(min_row=1, max_row=1), None)
    if not first_row:
        return header
    for cell in first_row:
        name = str(cell.value).strip() if cell.value is not None else ""
        if name:
            header[cell.column] = name
    return header


def _columns_in_sqref(sqref, header_map: dict) -> list[str]:
    """Header names covered by an openpyxl sqref range string, e.g. 'B2:B100'."""
    cols = set()
    for token in str(sqref).split():
        match = re.match(r"\$?([A-Z]+)\$?\d+(?::\$?([A-Z]+)\$?\d+)?", token)
        if not match:
            continue
        start_idx = column_index_from_string(match.group(1))
        end_idx = column_index_from_string(match.group(2) or match.group(1))
        cols.update(range(min(start_idx, end_idx), max(start_idx, end_idx) + 1))

    names = []
    for idx in sorted(cols):
        name = header_map.get(idx)
        if name and name not in names:
            names.append(name)
    return names


# ============================================================
# DATA VALIDATION -> column_rules (allowed_values / min-max / length)
# ============================================================

_SUPPORTED_DV_TYPES = ("list", "whole", "decimal", "textlength")


def extract_data_validation_rules(path: str) -> list[dict]:
    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        logger.exception("Could not open workbook for data validation extraction: %s", path)
        return []

    rules = []
    try:
        ws = wb.active
        header_map = _header_map(ws)

        for dv in ws.data_validations.dataValidation:
            dv_type = (dv.type or "").lower()
            if dv_type not in _SUPPORTED_DV_TYPES:
                continue

            columns = _columns_in_sqref(dv.sqref, header_map)
            if not columns:
                continue

            constraint = _constraint_for_validation(wb, ws, dv, dv_type)
            if not constraint:
                continue

            for header_name in columns:
                rules.append({"name": header_name, **constraint})
    finally:
        wb.close()

    return rules


def _to_number(formula):
    if formula is None:
        return None
    text = str(formula).strip().lstrip("=")
    try:
        return float(text)
    except ValueError:
        return None


_RANGE_MIN_MAX_OPERATORS = {
    "between": "between",
    "notbetween": None,  # not representable without a "forbidden range" concept — skip
    "greaterthan": "min",
    "greaterthanorequal": "min",
    "lessthan": "max",
    "lessthanorequal": "max",
    "equal": "exact",
}


def _range_constraint(dv, numeric: bool) -> dict | None:
    """Turn a whole/decimal/textLength validation's operator + formula(s)
    into {min_value,max_value} or {min_length,max_length}. Only a
    literal number in formula1/formula2 is supported — a formula
    referencing another cell is left alone rather than guessed at."""

    op = (dv.operator or "between").lower()
    kind = _RANGE_MIN_MAX_OPERATORS.get(op)
    if not kind:
        return None

    f1 = _to_number(dv.formula1)
    if f1 is None:
        return None

    min_key = "min_value" if numeric else "min_length"
    max_key = "max_value" if numeric else "max_length"

    constraints = {}
    if kind == "between":
        f2 = _to_number(dv.formula2)
        if f2 is None:
            return None
        constraints[min_key] = min(f1, f2)
        constraints[max_key] = max(f1, f2)
    elif kind == "min":
        constraints[min_key] = f1
    elif kind == "max":
        constraints[max_key] = f1
    elif kind == "exact":
        constraints[min_key] = f1
        constraints[max_key] = f1

    if not numeric:
        constraints = {k: int(v) for k, v in constraints.items()}

    return constraints or None


def _resolve_list_values(wb, ws, formula1) -> list[str]:
    """A dropdown's allowed values: either a literal '"A,B,C"' or a
    reference to a range of cells holding the list."""
    if not formula1:
        return []

    text = str(formula1).strip()

    if text.startswith('"') and text.endswith('"'):
        return [v.strip() for v in text[1:-1].split(",") if v.strip()]

    text = text.lstrip("=")
    match = re.match(
        r"(?:'?([^'!]+)'?!)?\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$",
        text,
    )
    if not match:
        return []

    sheet_name, c1, r1, c2, r2 = match.groups()
    target_ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else ws
    c2 = c2 or c1
    r2 = r2 or r1

    col1, col2 = column_index_from_string(c1), column_index_from_string(c2)
    row1, row2 = int(r1), int(r2)

    values = []
    for row in target_ws.iter_rows(
        min_row=min(row1, row2), max_row=max(row1, row2),
        min_col=min(col1, col2), max_col=max(col1, col2),
    ):
        for cell in row:
            if cell.value is not None and str(cell.value).strip():
                values.append(str(cell.value).strip())
    return values


def _constraint_for_validation(wb, ws, dv, dv_type: str) -> dict | None:
    if dv_type == "list":
        values = _resolve_list_values(wb, ws, dv.formula1)
        return {"allowed_values": values} if values else None

    if dv_type in ("whole", "decimal"):
        return _range_constraint(dv, numeric=True)

    if dv_type == "textlength":
        return _range_constraint(dv, numeric=False)

    return None


# ============================================================
# FORMULAS -> cross-column rule_type "formula" rules
# ============================================================

_CELL_REF_RE = re.compile(r"(?<![A-Za-z0-9_])\$?([A-Z]{1,3})\$?(\d+)(?![A-Za-z0-9_(])")
_FUNCTION_CALL_RE = re.compile(r"[A-Za-z_][A-Za-z0-9_.]*\s*\(")

# Only scan the first few data rows — enough to see every column's own
# template formula (assumed consistent down the column) without
# walking a large uploaded-as-a-template file cell by cell.
_FORMULA_SCAN_ROWS = 5


def extract_formula_rules(path: str) -> list[dict]:
    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        logger.exception("Could not open workbook for formula extraction: %s", path)
        return []

    rules = []
    try:
        ws = wb.active
        header_map = _header_map(ws)
        col_letter_to_header = {
            get_column_letter(idx): name for idx, name in header_map.items()
        }

        seen_columns = set()
        max_row = min(ws.max_row or 1, _FORMULA_SCAN_ROWS + 1)

        for row in ws.iter_rows(min_row=2, max_row=max_row):
            for cell in row:
                if cell.data_type != "f":
                    continue
                header_name = header_map.get(cell.column)
                if not header_name or header_name in seen_columns:
                    continue

                rule = _formula_rule(cell, header_name, col_letter_to_header)
                seen_columns.add(header_name)  # don't retry this column even if this row's formula didn't qualify
                if rule:
                    rules.append(rule)
    finally:
        wb.close()

    return rules


def _formula_rule(cell, header_name: str, col_letter_to_header: dict) -> dict | None:
    formula = cell.value
    if not isinstance(formula, str):
        return None

    body = formula.lstrip("=")

    if _FUNCTION_CALL_RE.search(body) or "!" in body:
        # Function calls (SUM, IF, ROUND, ...) and cross-sheet
        # references are out of scope for a safe arithmetic rule —
        # left alone rather than guessed at.
        return None

    self_row = cell.row
    refs = list(_CELL_REF_RE.finditer(body))
    if not refs:
        return None

    operand_columns = {}
    rewritten = body

    # Replace right-to-left so earlier matches' offsets (taken from the
    # original `body`) stay valid as later-in-string matches are spliced in.
    for match in reversed(refs):
        col_letters, row_num = match.group(1), int(match.group(2))
        if row_num != self_row:
            return None  # only same-row references are supported

        source_header = col_letter_to_header.get(col_letters)
        if not source_header or source_header == header_name:
            return None

        var_name = next(
            (v for v, h in operand_columns.items() if h == source_header),
            None,
        )
        if var_name is None:
            var_name = f"v{len(operand_columns) + 1}"
            operand_columns[var_name] = source_header

        rewritten = rewritten[: match.start()] + var_name + rewritten[match.end():]

    if not operand_columns or not safe_expr.is_safe_expression(rewritten):
        return None

    return {
        "name": header_name,
        "rule_type": "formula",
        "expression": rewritten,
        "operand_columns": operand_columns,
        "raw_formula": body,
        "tolerance": 0.01,
        "source": "template_formula",
    }


# ============================================================
# CONDITIONAL FORMATTING -> advisory notes (not enforced)
# ============================================================

_CF_OPERATOR_PHRASES = {
    "greaterThan": "is greater than",
    "greaterThanOrEqual": "is greater than or equal to",
    "lessThan": "is less than",
    "lessThanOrEqual": "is less than or equal to",
    "equal": "equals",
    "notEqual": "does not equal",
    "between": "is between",
    "notBetween": "is not between",
}


def extract_conditional_formatting_notes(path: str) -> list[str]:
    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        logger.exception("Could not open workbook for conditional formatting extraction: %s", path)
        return []

    notes = []
    try:
        ws = wb.active
        header_map = _header_map(ws)

        for cf_range in ws.conditional_formatting:
            columns = _columns_in_sqref(cf_range.sqref, header_map)
            label = ", ".join(f"'{c}'" for c in columns) if columns else "A range"

            for rule in cf_range.rules:
                sentence = _cf_sentence(label, rule)
                if sentence:
                    notes.append(sentence)
    finally:
        wb.close()

    return notes


def _cf_sentence(label: str, rule) -> str | None:
    if getattr(rule, "type", None) != "cellIs":
        return None

    phrase = _CF_OPERATOR_PHRASES.get(getattr(rule, "operator", None))
    if not phrase:
        return None

    formula = rule.formula or []
    values = ", ".join(str(f) for f in formula)
    if not values:
        return None

    return f"{label} is flagged when the value {phrase} {values}."


# ============================================================
# VBA MACROS -> raw source (best-effort, advisory only)
# ============================================================

def extract_vba_source(path: str) -> str:
    """
    Raw macro source from an .xlsm file's VBA project, concatenated
    and labeled by module.

    Uses oletools.olevba, a static macro-analysis library built for
    reading (never executing) potentially malicious VBA — the
    appropriate tool for a macro coming from an uploaded file. Returns
    "" for anything else (not .xlsm, no macros, oletools missing, or
    the file can't be parsed) rather than raising, since VBA is only
    ever advisory material here.
    """
    if not path.lower().endswith(".xlsm"):
        return ""

    if VBA_Parser is None:
        logger.warning("oletools is not installed — skipping VBA macro extraction for %s", path)
        return ""

    parser = None
    try:
        parser = VBA_Parser(path)
        if not parser.detect_vba_macros():
            return ""

        parts = []
        for _filename, stream_path, vba_filename, code in parser.extract_macros():
            code = (code or "").strip()
            if code:
                parts.append(f"--- Macro module: {vba_filename or stream_path} ---\n{code}")
        return "\n\n".join(parts)
    except Exception:
        logger.exception("Could not extract VBA macros from %s", path)
        return ""
    finally:
        if parser is not None:
            try:
                parser.close()
            except Exception:
                pass


# ============================================================
# ORCHESTRATOR
# ============================================================

def extract_excel_business_rules(path: str) -> dict:
    """
    Everything this module can mine from one Excel template:

        column_rules     deterministic — data validation + formula
                          rules, in validation_service's rule shape
        advisory_notes   conditional formatting, human-readable only
        vba_text         raw macro source (.xlsm only), for best-effort
                          interpretation elsewhere — never a hard rule
    """
    column_rules = extract_data_validation_rules(path) + extract_formula_rules(path)

    return {
        "column_rules": column_rules,
        "advisory_notes": extract_conditional_formatting_notes(path),
        "vba_text": extract_vba_source(path),
    }
