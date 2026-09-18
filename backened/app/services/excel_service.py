import openpyxl
from fastapi import HTTPException, status


def extract_column_headers(file_path: str) -> list[str]:
    """
    Opens an Excel file and reads the header row (first row) of the first
    sheet, returning a clean list of column names.
    """
    try:
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not read the Excel file. Make sure it's a valid .xlsx file. ({e})",
        )

    sheet = workbook.active
    first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    workbook.close()

    if not first_row:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file appears to be empty.",
        )

    headers = [str(cell).strip() for cell in first_row if cell is not None and str(cell).strip()]

    if not headers:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No column headers were found in the first row of the file.",
        )

    return headers


def read_data_rows(file_path: str) -> list[dict]:
    """
    Reads every data row (everything after the header row) from an Excel
    file and returns it as a list of dicts: [{"Column Name": value, ...}, ...]
    Row numbers are 1-based and refer to the actual Excel row (so row 2 is
    the first data row, since row 1 is the header) — this makes error
    messages like "row 5" line up with what the user sees when they open
    the file themselves.
    """
    try:
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not read the Excel file. ({e})",
        )

    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)

    header_row = next(rows_iter, None)
    if not header_row:
        workbook.close()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="The uploaded file appears to be empty.")

    headers = [str(cell).strip() if cell is not None else "" for cell in header_row]

    data_rows = []
    for row_index, row_values in enumerate(rows_iter, start=2):  # start=2: Excel row numbering after header
        # Skip fully blank rows
        if all(v is None or str(v).strip() == "" for v in row_values):
            continue

        row_dict = {"_row_number": row_index}
        for col_name, value in zip(headers, row_values):
            if col_name:
                row_dict[col_name] = value
        data_rows.append(row_dict)

    workbook.close()
    return data_rows


def read_file_preview(
    file_path: str,
    max_rows: int = 25,
    max_chars: int = 6000,
) -> str:
    """
    Builds a compact, human-readable text preview of an uploaded Excel file
    for the AI assistant: the column headers, the total number of data rows,
    and the first few rows so the model can reason about actual cell values
    instead of only header names.

    Never raises — a malformed file simply yields a short "could not read"
    note, so a chat turn can always fall back to the data it already has.
    """
    try:
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        return f"(Could not read the file for preview: {e})"

    try:
        try:
            sheet = workbook.active
            rows_iter = sheet.iter_rows(values_only=True)

            header_row = next(rows_iter, None)
            if not header_row:
                return "(The uploaded file appears to be empty.)"

            headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
            used_headers = [header for header in headers if header]

            preview_rows = []
            total_data_rows = 0
            for row_values in rows_iter:
                if all(value is None or str(value).strip() == "" for value in row_values):
                    continue
                total_data_rows += 1
                if len(preview_rows) < max_rows:
                    preview_rows.append(row_values)
        except Exception as e:
            return f"(Could not read the file for preview: {e})"
    finally:
        workbook.close()

    lines = [
        "COLUMN HEADERS:",
        ", ".join(used_headers) or "(none)",
        "",
        f"TOTAL DATA ROWS: {total_data_rows}",
        "",
        f"FIRST {len(preview_rows)} ROWS:",
    ]

    for index, row_values in enumerate(preview_rows, start=1):
        cells = []
        for header, value in zip(headers, row_values):
            if not header:
                continue
            value_text = "" if value is None else str(value).strip()
            cells.append(f"{header}={value_text}")
        lines.append(f"Row {index}: " + " | ".join(cells))

    text = "\n".join(lines)
    if len(text) > max_chars:
        text = text[:max_chars] + "\n...(preview truncated)"

    return text